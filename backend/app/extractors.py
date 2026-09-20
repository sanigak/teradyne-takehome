"""Deterministic extraction with original-source locations and honest attribution.

Office parsers are intentionally local.  Legacy binary formats are normalized by
LibreOffice in an isolated temporary profile; originals are never overwritten.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import zipfile


SUPPORTED_EXTENSIONS = frozenset({".md", ".doc", ".docx", ".ppt", ".pptx", ".xls", ".xlsx"})
LEGACY_TARGETS = {".doc": "docx", ".ppt": "pptx", ".xls": "xlsx"}
_DEFAULT_AUTHORS = {"", "unknown", "unknown author", "author", "python-docx", "python-pptx", "openpyxl", "microsoft office user", "microsoft office", "microsoft word", "microsoft excel", "microsoft powerpoint", "libreoffice", "user"}
MAX_SOURCE_BYTES = 32 * 1024 * 1024
MAX_OFFICE_EXPANDED_BYTES = 64 * 1024 * 1024


@dataclass
class ExtractedChunk:
    text: str
    locator: str


@dataclass
class ExtractedDocument:
    title: str
    author: str | None
    attendees: list[str]
    date: str | None
    chunks: list[ExtractedChunk]
    warnings: list[str]


class ExtractionError(ValueError):
    """A per-file failure which ingestion can report without aborting the batch."""


def discover_soffice(configured: str | None = None) -> str:
    candidate = configured or os.environ.get("SOFFICE_PATH")
    if candidate:
        resolved = shutil.which(candidate) or (str(Path(candidate).resolve()) if Path(candidate).is_file() else None)
        if resolved:
            return resolved
        raise ExtractionError("Configured SOFFICE_PATH does not identify a LibreOffice executable.")
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for root in (os.environ.get("PROGRAMFILES", "C:/Program Files"), os.environ.get("PROGRAMFILES(X86)", "C:/Program Files (x86)")):
        found = Path(root) / "LibreOffice/program/soffice.exe"
        if found.is_file():
            return str(found)
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if mac.is_file():
        return str(mac)
    raise ExtractionError("Legacy Office extraction requires LibreOffice. Install it and set SOFFICE_PATH to its soffice executable.")


def _clean(value: object) -> str:
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip() if value is not None else ""


def _append(chunks: list[ExtractedChunk], text: object, locator: str) -> None:
    content = _clean(text)
    if content:
        chunks.append(ExtractedChunk(content, locator))


def _markdown(path: Path) -> tuple[list[ExtractedChunk], str | None, str | None, list[str]]:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    chunks: list[ExtractedChunk] = []
    start = 0
    block: list[str] = []
    for index, line in enumerate(lines + [""]):
        if line.strip():
            if not block:
                start = index + 1
            block.append(line)
        elif block:
            _append(chunks, "\n".join(block), f"Lines {start}-{index}")
            block = []
    title = next((line.removeprefix("# ").strip() for line in lines if line.startswith("# ")), None)
    return chunks, title, None, []


def _word(path: Path) -> tuple[list[ExtractedChunk], str | None, str | None, list[str]]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    chunks: list[ExtractedChunk] = []
    paragraph = table = 0
    for element in document.element.body:
        if element.tag == qn("w:p"):
            paragraph += 1
            _append(chunks, Paragraph(element, document).text, f"Paragraph {paragraph}")
        elif element.tag == qn("w:tbl"):
            table += 1
            value = Table(element, document)
            for row_index, row in enumerate(value.rows, 1):
                cells = [cell.text.strip() for cell in row.cells]
                _append(chunks, " | ".join(f"Cell {i}: {text}" for i, text in enumerate(cells, 1) if text), f"Table {table}, row {row_index}, cells 1-{len(cells)}")
    props = document.core_properties
    return chunks, props.title or None, props.author or None, []


def _slides(path: Path) -> tuple[list[ExtractedChunk], str | None, str | None, list[str]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    chunks: list[ExtractedChunk] = []

    def shape_chunks(shapes: object, prefix: str) -> None:
        for shape_index, shape in enumerate(shapes, 1):
            locator = f"{prefix}, shape {shape_index}"
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                shape_chunks(shape.shapes, f"{locator}, group")
            if shape.has_text_frame:
                _append(chunks, shape.text_frame.text, locator)
            if shape.has_table:
                for row_index, row in enumerate(shape.table.rows, 1):
                    _append(chunks, " | ".join(f"Cell {i}: {cell.text}" for i, cell in enumerate(row.cells, 1) if cell.text.strip()), f"{locator}, table row {row_index}")

    for index, slide in enumerate(presentation.slides, 1):
        shape_chunks(slide.shapes, f"Slide {index}")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None:
                _append(chunks, notes.text, f"Slide {index}, speaker notes")
    props = presentation.core_properties
    title = props.title or None
    if presentation.slides and (not title or title.strip().casefold() == "powerpoint presentation"):
        first_slide = presentation.slides[0]
        title_shape = first_slide.shapes.title
        if title_shape is not None and title_shape.text.strip():
            title = title_shape.text.strip()
        else:
            # Some LibreOffice versions discard legacy PPT core properties and
            # convert title textboxes to ordinary auto-shapes. Recognize a
            # visibly styled title followed immediately by the metadata block;
            # arbitrary opening body text must not be skipped as a title.
            text_shapes = [shape for shape in first_slide.shapes if shape.has_text_frame and shape.text.strip()]
            if len(text_shapes) >= 2:
                heading, header = text_shapes[:2]
                sizes = [font.size.pt for paragraph in heading.text_frame.paragraphs
                         for font in [paragraph.font, *(run.font for run in paragraph.runs)] if font.size is not None]
                if ("\n" not in heading.text.strip() and len(heading.text.strip()) <= 250
                        and heading.top < header.top and heading.top < presentation.slide_height / 3
                        and sizes and max(sizes) >= 24
                        and re.match(r"^(?:Title|Author|Date|Attendees)\s*:", header.text.strip(), re.I)):
                    title = heading.text.strip()
    return chunks, title, props.author or None, []


def _sheets(path: Path) -> tuple[list[ExtractedChunk], str | None, str | None, list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=False, read_only=True)
    cached = load_workbook(path, data_only=True, read_only=True)
    chunks: list[ExtractedChunk] = []
    warnings: list[str] = []
    try:
        for sheet in workbook.worksheets:
            # OOXML can declare an enormous used range with only a few bytes of
            # XML. Refuse it before iter_rows allocates/visits that range.
            if (sheet.max_row or 0) * (sheet.max_column or 0) > 200_000:
                raise ExtractionError(f"Sheet '{sheet.title}' exceeds the 200,000-cell extraction limit; remove unused formatting or split the workbook.")
            # Stream the formula and cached-value views together. Random cell
            # lookups in read-only worksheets repeatedly reparse earlier rows.
            for row, cached_row in zip(sheet.iter_rows(), cached[sheet.title].iter_rows()):
                values: list[str] = []
                coordinates: list[str] = []
                for cell, cached_cell in zip(row, cached_row):
                    if cell.value is None:
                        continue
                    value = cell.value
                    if cell.data_type == "f":
                        saved_value = cached_cell.value
                        if saved_value is None:
                            warnings.append(f"{sheet.title}!{cell.coordinate}: formula has no cached value; formula text retained and not evaluated.")
                            value = f"{value} [cached result unavailable]"
                        else:
                            value = f"{saved_value} [formula: {value}; cached value may be stale]"
                            warnings.append(f"{sheet.title}!{cell.coordinate}: using saved formula result, which may be stale.")
                    values.append(f"{cell.coordinate}: {_clean(value)}")
                    coordinates.append(cell.coordinate)
                if coordinates:
                    _append(chunks, " | ".join(values), f"Sheet '{sheet.title}', cells {coordinates[0]}:{coordinates[-1]}")
        return chunks, workbook.properties.title or None, workbook.properties.creator or None, warnings
    finally:
        workbook.close()
        cached.close()


def _metadata(chunks: list[ExtractedChunk], title: str | None, core_author: str | None, stem: str, warnings: list[str]) -> ExtractedDocument:
    # Only the leading metadata block is attribution. A quoted Author: line in
    # meeting dialogue, tables, notes, or a later slide cannot become an expert.
    # Cell-address labels allow the same visible convention in spreadsheets.
    header_chunks = chunks[:30]
    if chunks and chunks[0].locator.startswith("Slide "):
        header_chunks = [chunk for chunk in header_chunks if chunk.locator.startswith("Slide 1,") and "speaker notes" not in chunk.locator]
    source = "\n".join(chunk.text for chunk in header_chunks)
    source = re.sub(r"(?m)^\$?[A-Z]{1,3}\$?\d+:\s*", "", source)
    fields: dict[str, str] = {}
    normalize_title = lambda value: re.sub(r"[^\w]", "", value or "").casefold()
    for index, line in enumerate(source.splitlines()):
        if not line.strip():
            continue
        match = re.fullmatch(r"(?:\*\*)?(Title|Author|Date|Attendees|Client|Domain|Priority|Classification)(?:\*\*)?\s*:\s*([^|]+)", line.strip(), flags=re.I)
        if match:
            key, value = match.group(1).lower(), match.group(2).strip().strip("*")
            if key in fields and fields[key] != value:
                warnings.append(f"Conflicting {key} values in source metadata; first header value retained.")
            fields.setdefault(key, value)
        elif index == 0 and (line.startswith("# ") or (title and normalize_title(line) == normalize_title(title))):
            continue
        else:
            break
    author = fields.get("author") or core_author
    if author and author.strip().lower() in _DEFAULT_AUTHORS:
        author = None
    attendees = list(dict.fromkeys(person.strip() for person in re.split(r"[;,]", fields.get("attendees", "")) if person.strip()))
    source_date = fields.get("date")
    if source_date:
        try:
            source_date = date.fromisoformat(source_date).isoformat()
        except ValueError:
            warnings.append("Source date is not a valid ISO date (YYYY-MM-DD); date left unset.")
            source_date = None
    if not author:
        warnings.append("Source author is missing; no author was inferred.")
    if not source_date:
        warnings.append("Source date is missing; filesystem/Office creation times are not treated as meeting dates.")
    if not chunks:
        raise ExtractionError("No extractable text found. Image-only documents and OCR are unsupported.")
    return ExtractedDocument(fields.get("title") or title or stem, author, attendees, source_date, chunks, warnings)


def extract_document(path: Path, *, soffice_path: str | None = None) -> ExtractedDocument:
    """Read one supported source, retaining stable locators in its original format."""
    path = Path(path)
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(f"Unsupported source extension: {extension or '(none)'}")
    if not path.is_file():
        raise ExtractionError(f"Source file does not exist: {path.name}")
    try:
        if path.stat().st_size > MAX_SOURCE_BYTES:
            raise ExtractionError("Source exceeds the 32 MiB extraction limit; split the source before ingestion.")
        if extension in {".docx", ".pptx", ".xlsx"}:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 10_000 or sum(entry.file_size for entry in entries) > MAX_OFFICE_EXPANDED_BYTES:
                    raise ExtractionError("Office archive exceeds extraction limits (10,000 entries / 64 MiB expanded); split the source before ingestion.")
        if extension in LEGACY_TARGETS:
            executable = discover_soffice(soffice_path)
            target = LEGACY_TARGETS[extension]
            with tempfile.TemporaryDirectory(prefix="relay-extract-") as temporary:
                root = Path(temporary)
                profile = root / "profile"
                output = root / "converted"
                output.mkdir()
                # A unique profile prevents a GUI instance or simultaneous ingestion
                # from swallowing a headless conversion invocation.
                result = subprocess.run([executable, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", target, "--outdir", str(output), str(path.resolve())], capture_output=True, text=True, timeout=90, check=False)
                converted = output / f"{path.stem}.{target}"
                if result.returncode or not converted.is_file() or converted.stat().st_size == 0:
                    raise ExtractionError(f"LibreOffice could not convert {path.name}; exit code {result.returncode}. Verify that the source is a valid, unencrypted Office file.")
                document = extract_document(converted)
                document.warnings.insert(0, f"Original {extension} source normalized with LibreOffice; locators describe converted document structure, not printed page numbers.")
                return document
        chunks, title, author, warnings = {".md": _markdown, ".docx": _word, ".pptx": _slides, ".xlsx": _sheets}[extension](path)
        return _metadata(chunks, title, author, path.stem, warnings)
    except ExtractionError:
        raise
    except subprocess.TimeoutExpired as exc:
        raise ExtractionError(f"LibreOffice conversion timed out after 90 seconds for {path.name}.") from exc
    except Exception as exc:
        raise ExtractionError(f"Cannot extract {path.name}: {type(exc).__name__}. Check file integrity and installed Office parser dependencies.") from exc
