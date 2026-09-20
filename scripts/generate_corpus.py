"""Reproduce the fictional Relay AI dataset, including real legacy Office files.

Usage: python scripts/generate_corpus.py [--soffice PATH] [--output data/corpus]
Specifications live in this script, outside the ingested corpus. No source file is
a renamed modern-format file: LibreOffice exports DOC/PPT/XLS binary formats.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from backend.app.extractors import discover_soffice, extract_document


PEOPLE = {
    "maya": "Maya Chen", "theo": "Theo Haddad", "sofia": "Sofia Ortiz", "liam": "Liam Okafor",
    "priya": "Priya Raman", "noah": "Noah Brooks", "emma": "Emma Laurent", "jules": "Jules Park",
    "ines": "Ines Duarte", "oscar": "Oscar Bell", "nina": "Nina Shah", "elliot": "Elliot Reed",
}


def spec(filename: str, title: str, client: str, day: int, author: str, people: str, domain: str, priority: str, blocks: list[tuple[str, str]], facts: list[str], *, notes: str = "", table: list[list[object]] | None = None) -> dict:
    return dict(filename=filename, title=title, client=client, date=f"2026-09-{day:02d}", author=PEOPLE[author], attendees=[PEOPLE[p] for p in people.split()], domain=domain, priority=priority, blocks=blocks, expected_facts=facts, notes=notes, table=table)


SPECS = [
    spec("atlas_01_kickoff.md", "Atlas Forge — deployment kickoff", "Atlas Forge", 1, "maya", "maya theo sofia liam", "Deployment planning", "High", [
        ("Maya Chen", "We are building an internal maintenance-manual assistant for Atlas Forge technicians. This pilot can recommend a cited procedure, but it cannot command a machine or close a work order."),
        ("Liam Okafor", "The service team has 400 manuals and 22 volunteer technicians. Manuals contain serial numbers but should contain no customer contact details. I own the approved manual list and technician onboarding."),
        ("Sofia Ortiz", "I propose a private gateway in the client's existing VPC. We will use the Frankfurt region and retain tenant identifiers on every indexed chunk. The deployment team still needs permission to create the private endpoint."),
        ("Maya Chen", "Planning assumption, not a signed commitment: target October 15, 2026 for the pilot launch. The date assumes the security review closes in September. I will update this note if that dependency changes."),
        ("Theo Haddad", "I own the Atlas security sign-off. Do not send raw telemetry or operator names to the external model provider. We need a written redaction test and an access review before production use."),
        ("Liam Okafor", "Decision: the first pilot includes only maintenance manuals, not incident tickets. That keeps the initial data scope manageable. An unanswered question should go to an accountable manual owner."),
        ("Maya Chen", "Actions: Liam supplies the manual inventory by September 4. Sofia documents the gateway design by September 7. Theo schedules the security review for September 10."),
        ("Sofia Ortiz", "Open point: the client's identity team has not yet assigned the service principal. I will track that as a blocker, not silently substitute a shared account."),
    ], ["Initial October 15 launch is provisional.", "Atlas pilot is advisory and uses maintenance manuals only."]),
    spec("atlas_02_architecture.md", "Atlas Forge — private gateway architecture", "Atlas Forge", 7, "sofia", "sofia theo liam", "Architecture", "High", [
        ("Sofia Ortiz", "The Atlas gateway design is ready. The application runs inside the Atlas private VPC in Frankfurt. Requests carry an authenticated tenant claim, and the retriever filters chunks by that claim before ranking."),
        ("Theo Haddad", "Decision: model-bound requests must remove operator names, email addresses, and raw device telemetry. A redaction failure must block the request. No best-effort bypass is permitted."),
        ("Liam Okafor", "What does a technician see when redaction blocks a request? The maintenance workflow must stay understandable."),
        ("Sofia Ortiz", "They see a request-blocked explanation and can rephrase the question. The system will retain only a diagnostic code and request ID in application logs. It will not log the original prompt."),
        ("Theo Haddad", "The architecture decision is approved for the pilot, subject to the security gate. I still need the least-privilege role matrix and proof that the redaction tests include serial-number variants."),
        ("Sofia Ortiz", "Action: I will deliver the role matrix by September 9. The private endpoint request is pending with Atlas identity engineering. Liam can coordinate that team, but I own the technical integration."),
        ("Liam Okafor", "We should include a five-percent canary in the release plan. That is a proposal today; Maya will confirm the release sequence with operations."),
    ], ["Atlas runs in a private VPC in Frankfurt.", "Redaction failure blocks a model request."]),
    spec("atlas_03_security_gate.md", "Atlas Forge — security gate and revised launch", "Atlas Forge", 10, "theo", "theo maya sofia", "Security and release", "Critical", [
        ("Theo Haddad", "The security review found that two operator-name formats escaped redaction. We will not approve the October pilot while those cases fail. Sofia owns the fix and I own the approval decision."),
        ("Maya Chen", "Decision: the Atlas pilot launch moves to November 2, 2026. This explicitly supersedes the provisional October 15, 2026 target recorded at kickoff on September 1."),
        ("Sofia Ortiz", "I can land the redaction patch by September 18 and attach replay evidence. The issue is in normalization before policy evaluation, not the tenant filter."),
        ("Theo Haddad", "Release gate: all redaction test cases must pass, the least-privilege role matrix must be approved, and the client security owner must acknowledge the runbook. I will record sign-off only after all three conditions hold."),
        ("Maya Chen", "Action: I will tell the Atlas sponsor that November 2 is the agreed launch date, conditional on this gate. We are not accepting extra incident-ticket ingestion in the current pilot."),
        ("Sofia Ortiz", "Telemetry retention details belong in the security addendum. We should not confuse operational diagnostics with raw prompts; the latter are prohibited in logs."),
        ("Theo Haddad", "Open question: nobody has confirmed the provider's zero-retention contractual addendum. I will request it; do not claim that agreement exists."),
    ], ["November 2, 2026 explicitly supersedes provisional October 15.", "Theo Haddad owns security approval."]),
    spec("atlas_04_delivery.md", "Atlas Forge — delivery checkpoint", "Atlas Forge", 14, "maya", "maya sofia liam", "Delivery operations", "High", [
        ("Maya Chen", "Checkpoint for Atlas: the pilot remains November 2, 2026 under the revised security decision. Redaction replay evidence and the identity service principal are the two open launch blockers."),
        ("Sofia Ortiz", "The service principal ticket is still pending. I can test the gateway locally but cannot run the production private-endpoint verification until Atlas identity engineering grants the scoped role."),
        ("Liam Okafor", "I will escalate the service principal ticket with Atlas identity engineering by September 16. Sofia remains the owner of the gateway and integration tests."),
        ("Maya Chen", "Decision: approve a five-percent canary for one business day, then twenty-five percent for one business day, then full pilot traffic only if latency and error budgets pass. The release deck will carry the exact stop conditions."),
        ("Sofia Ortiz", "The latency workbook sets a p95 target below 2.5 seconds. Do not turn an estimated budget into a measured result; end-to-end production timings are not available yet."),
        ("Liam Okafor", "Do we know the permitted maintenance window for a rollback? I have not seen that agreed by the factory team."),
        ("Maya Chen", "No approved rollback maintenance window exists in our current documents. Route that question to me and Liam, who coordinates the factory owner. Action: request a written window at the operations review."),
    ], ["Atlas canary is 5%, then 25%, then full pilot traffic.", "Rollback maintenance window is undocumented."]),
    spec("beacon_01_discovery.md", "Beacon Route — forecasting discovery", "Beacon Route", 2, "priya", "priya noah emma jules", "Product discovery", "Medium", [
        ("Priya Raman", "Beacon wants a next-day depot volume forecast so planners can allocate staff. The first pilot is a daily planning aid for the North and Central depots. It will not automatically dispatch vehicles."),
        ("Jules Park", "The current baseline is a seasonal spreadsheet. Planners need the forecast by 06:00 local depot time and a visible reason when data is incomplete."),
        ("Noah Brooks", "I own the data connector and model evaluation. We need at least twelve weeks of shipment counts, weather flags, and depot closures; individual driver names are outside scope."),
        ("Emma Laurent", "I own commercial approval. The early cost estimate is not a signed price. I will negotiate the pilot statement of work once Priya freezes the two-depot scope."),
        ("Priya Raman", "Decision: ship a batch forecast with manual planner approval. No live dispatch integration belongs in this pilot. A planner can reject a forecast and retain the existing staffing plan."),
        ("Noah Brooks", "Action: I will profile the shipment feed by September 5. We have no confirmed retention terms for the external weather vendor yet."),
        ("Jules Park", "I can supply closure calendars, but holiday annotations are incomplete. Mark that as a known data quality issue rather than fill missing holidays with guessed labels."),
    ], ["Beacon is batch staffing forecasting for two depots, not live dispatch."]),
    spec("beacon_02_data_review.md", "Beacon Route — connector and vendor data review", "Beacon Route", 5, "noah", "noah priya emma", "Data engineering", "High", [
        ("Noah Brooks", "The shipment connector is healthy for North depot, but Central has duplicate shipment IDs after retries. I counted 37 duplicates in the sample. Deduplicate on shipment ID plus depot, keeping the most recent update timestamp."),
        ("Priya Raman", "Decision: block the Central backtest until deduplication is verified. North can continue through offline evaluation. This is a data acceptance gate, not a reason to invent missing records."),
        ("Noah Brooks", "Action: I will add the composite-key test and produce a clean twelve-week sample by September 9. Jules owns the depot closure calendar and will validate the holiday flags."),
        ("Emma Laurent", "The weather vendor has not answered the retention and data-processing questions. I own the contract request. There is no signed data-processing addendum and no documented deletion SLA yet."),
        ("Priya Raman", "We can evaluate with a static weather extract while contracts are unresolved. Do not activate the ongoing vendor feed until Emma confirms the agreement."),
        ("Noah Brooks", "I will log row counts and freshness checks in the quality register. The register should distinguish a threshold from an observed measurement so reviewers cannot confuse a target with a passing result."),
    ], ["Central duplicate key is shipment ID plus depot, latest update wins.", "Weather vendor deletion SLA is unknown; Emma Laurent owns follow-up."]),
    spec("beacon_03_pilot_acceptance.md", "Beacon Route — pilot acceptance decision", "Beacon Route", 11, "priya", "priya noah jules", "Model evaluation", "High", [
        ("Noah Brooks", "After deduplication, the held-out forecast error is 11.4 percent on the two-depot sample. The historical seasonal spreadsheet baseline was 16 percent. These are percentage MAE results, not classification accuracy."),
        ("Jules Park", "Planners can work with that if the run arrives on time. A perfect model after the morning staffing decision is useless."),
        ("Priya Raman", "Decision: accept the Beacon pilot only when mean absolute percentage error is at most 12 percent, forecasts arrive by 06:00 local depot time, and missing-input runs fall back to the seasonal baseline with a warning."),
        ("Noah Brooks", "I own measuring forecast error and freshness. The evaluation set is frozen at eight holdout weeks; training must not include those weeks. We have not measured live operational savings yet."),
        ("Jules Park", "Decision: every forecast requires manual planner approval. An abstention or fallback should explain which inputs are missing and name the depot."),
        ("Priya Raman", "Action: Noah posts the acceptance workbook before September 13. Jules leads five morning rehearsals. Commercial signature remains separate from technical acceptance."),
    ], ["Beacon target is error at most 12%; measured sample error 11.4%; baseline 16%.", "Forecasts due 06:00 local depot time."]),
    spec("beacon_04_commercial.md", "Beacon Route — signed pilot commercial decision", "Beacon Route", 15, "emma", "emma priya jules", "Commercial decisions", "High", [
        ("Emma Laurent", "The two-depot Beacon pilot statement of work is signed. The fixed pilot fee is USD 48,000 excluding tax. Payment is fifty percent at kickoff and fifty percent after the acceptance review."),
        ("Priya Raman", "Please call out why the earlier document says USD 62,000. I do not want the sponsor seeing two prices with no explanation."),
        ("Emma Laurent", "The September 4 commercial terms document is an unsigned estimate for three depots at USD 62,000. Today's signed USD 48,000 agreement covers two depots and supersedes that estimate for the pilot."),
        ("Jules Park", "Decision: North and Central are the only pilot depots. South depot and automatic dispatch remain excluded. A third depot needs a priced change request."),
        ("Priya Raman", "Two approved records currently disagree about model acceptance: our September 11 decision says at most 12 percent error, while the September 15 planner handoff says at most 10 percent. No reconciled decision exists yet. I own resolving this discrepancy before pilot acceptance. Morning delivery by 06:00 is unchanged."),
        ("Emma Laurent", "Action: I will circulate the signed scope summary today. The weather vendor deletion SLA is still unanswered and should be routed to me for vendor clarification; the signed pilot price does not settle that contract question."),
    ], ["Signed pilot price USD 48,000 for two depots supersedes unsigned USD 62,000 estimate.", "Approved error thresholds 12 percent versus 10 percent conflict; Priya Raman must reconcile."]),
    spec("cedar_01_kickoff.md", "Cedar Vale — policy assistant kickoff", "Cedar Vale", 3, "ines", "ines oscar nina elliot", "Product and privacy", "High", [
        ("Ines Duarte", "Cedar Vale is piloting an internal policy-search assistant for clinic administrators. It answers questions about staff policies and cites the approved handbook. It must not diagnose patients or recommend treatment."),
        ("Nina Shah", "I own policy content approval. Staff may paste accidental patient identifiers into questions, so we need a clear input warning and reliable redaction. The corpus itself must exclude patient records."),
        ("Oscar Bell", "I own access-control implementation. Each chunk will carry the clinic and policy access group. Filtering occurs before model context assembly, not after an answer is generated."),
        ("Elliot Reed", "I own evaluation and incident triage. We should test one user's inability to see another clinic's restricted policies and require citations on each policy claim."),
        ("Ines Duarte", "Decision: pilot only with administrative users in East and Harbor clinics. Escalate missing policy answers to Nina; do not generate a substitute policy."),
        ("Nina Shah", "Retention is still under review. The draft privacy review will propose a period, but no one should represent that draft as a final decision."),
    ], ["Cedar is administrative policy search, not patient diagnosis.", "Nina Shah owns policy approval."]),
    spec("cedar_02_access_review.md", "Cedar Vale — tenant isolation review", "Cedar Vale", 8, "oscar", "oscar elliot nina", "Identity and access", "Critical", [
        ("Oscar Bell", "Cedar now applies clinic_id and policy_group filters before retrieval ranking. Both claims must come from the authenticated identity token; the user cannot select an arbitrary clinic in a query parameter."),
        ("Elliot Reed", "I tested an East-clinic user against Harbor-only policy text. The restricted chunks never reached model context. We also need a missing-claim test."),
        ("Oscar Bell", "Decision: missing clinic or policy-group claims fail closed with a forbidden response. We do not fall back to a global collection. I own the middleware and the access regression tests."),
        ("Nina Shah", "Policy publication remains a human approval process. A user correction goes into review, not directly into the policy corpus. Otherwise a well-intentioned correction could become an unapproved rule."),
        ("Elliot Reed", "Action: I will add an identity-expiration case and a revoked-group case by September 12. Every evaluation record should keep the tested document version and retrieval evidence."),
        ("Oscar Bell", "An on-call weekend rota is not documented yet. I can discuss access failures but cannot promise a staffed weekend response time."),
    ], ["Cedar filters clinic_id and policy_group before ranking and fails closed on missing claims."]),
    spec("cedar_03_validation.md", "Cedar Vale — validation and privacy decision", "Cedar Vale", 13, "nina", "nina ines elliot oscar", "Evaluation and privacy", "High", [
        ("Elliot Reed", "The frozen policy test set contains 80 questions: 60 answerable policy questions, 12 deliberately missing policies, and 8 cross-clinic access attempts. We scored claim support separately from retrieval coverage."),
        ("Nina Shah", "Decision: delete raw user prompts within 24 hours. This supersedes the seven-day proposal in the September 5 privacy review. Keep only redacted operational counters for 30 days; no patient identifiers belong in analytics."),
        ("Ines Duarte", "To confirm, the seven-day period was a draft, not an active production policy. Nina owns this final privacy decision and Oscar will implement the scheduled deletion."),
        ("Oscar Bell", "I will add deletion job monitoring and verify a record older than 24 hours cannot be retrieved from the raw-prompt store. This retention rule is separate from the signed source-policy documents."),
        ("Elliot Reed", "Acceptance requires zero cross-clinic disclosures, at least 95 percent supported claims on answerable questions, and routing for deliberately absent policies. The score workbook records observed results and remaining review items."),
        ("Nina Shah", "Open issue: paid parental leave eligibility for temporary contractors is not defined in our approved handbook extracts. Route that policy question to me rather than borrowing rules from another employer."),
    ], ["Raw prompt deletion within 24 hours supersedes seven-day draft.", "Parental leave for temporary contractors is undocumented; route to Nina Shah."]),
    spec("cedar_04_incident_drill.md", "Cedar Vale — incident tabletop notes", "Cedar Vale", 16, "elliot", "elliot oscar ines", "Incident operations", "Critical", [
        ("Elliot Reed", "Tabletop, not a real incident: a stale role cache lets a revoked administrator request restricted snippets. The runbook severity is P1 for suspected cross-clinic disclosure."),
        ("Oscar Bell", "Decision: disable the Cedar assistant endpoint, revoke the affected role cache, and preserve redacted audit evidence. Do not copy raw patient identifiers into an incident chat."),
        ("Ines Duarte", "I own client incident communication. Oscar owns containment and access restoration; Elliot owns triage and the evidence timeline. Nina approves any policy-content changes after the incident."),
        ("Elliot Reed", "The drill took nine minutes to disable the endpoint. That is a rehearsal observation, not a contractual response-time guarantee. No signed weekend support SLA exists in this corpus."),
        ("Oscar Bell", "Restore only after the revoked-group regression test passes and Theo's independent security review is complete. Theo is a reviewer on the runbook, although not an attendee at today's drill."),
        ("Ines Duarte", "Action: Elliot updates the incident worksheet; Oscar attaches cache-invalidation evidence. If a user asks who guarantees a fifteen-minute weekend response, route to me for a commercial answer instead of inventing coverage."),
    ], ["Cedar disclosure response disables endpoint and revokes role cache.", "No signed weekend support SLA exists."]),
    spec("atlas_security_addendum.docx", "Atlas Forge — security logging addendum", "Atlas Forge", 12, "theo", "theo sofia", "Security", "Critical", [
        ("Approved logging boundaries", "Atlas diagnostic logs contain request IDs, error codes, duration, and redaction outcome only. Raw prompts, operator names, email addresses, and raw device telemetry must never be written to application logs."),
        ("Retention and location", "Retain redacted diagnostic logs for 30 days in Frankfurt. Theo Haddad owns approval of any retention change. The deletion job runs daily and Sofia Ortiz owns its operational alert."),
        ("Security gate", "Launch remains conditional on all redaction tests passing, an approved least-privilege role matrix, and client acknowledgment of the runbook. This addendum does not attest that the open provider zero-retention contract is signed."),
    ], ["Atlas redacted diagnostics retained 30 days in Frankfurt."]),
    spec("atlas_gateway_design.doc", "Atlas Forge — gateway role matrix", "Atlas Forge", 9, "sofia", "sofia theo liam", "Architecture", "High", [
        ("Identity design", "The Atlas gateway uses a workload service principal with a read-only manual-index scope. It must not grant write access to source manuals or use a technician's personal credential."),
        ("Failure handling", "A missing or mismatched tenant claim denies retrieval before the model is called. A redaction failure blocks the request. A private endpoint outage shows an availability error and must not fall back to an unapproved public endpoint."),
        ("Ownership and status", "Sofia Ortiz owns gateway integration. Liam Okafor coordinates Atlas identity engineering. The requested production service principal is pending; this design document is not proof of provisioning."),
    ], ["Gateway principal is read-only for manual index.", "No public-endpoint fallback."]),
    spec("atlas_release_plan.pptx", "Atlas Forge — staged release runbook", "Atlas Forge", 14, "maya", "maya sofia liam", "Deployment operations", "High", [
        ("Release sequence", "Launch target: November 2, 2026 after security sign-off. Start with 5 percent of pilot users for one business day, then 25 percent for one business day, then 100 percent only while release gates pass."),
        ("Stop conditions", "Pause rollout when p95 end-to-end latency is 2.5 seconds or greater, request error rate exceeds 1 percent over a 15-minute window, or any unauthorized retrieval is detected. Unauthorized retrieval triggers an immediate stop regardless of traffic volume."),
        ("Rollback ownership", "Sofia Ortiz owns technical rollback to the prior gateway version. Maya Chen owns sponsor communication. Liam Okafor coordinates the factory operating window, which is not yet approved."),
    ], ["Atlas rollout stops at p95 >=2.5 seconds or error rate >1% over 15 minutes."], notes="The canary percentages describe a planned release, not observed production adoption. Do not claim a rollback maintenance window has been agreed."),
    spec("atlas_latency_budget.xlsx", "Atlas Forge — latency and release budget", "Atlas Forge", 14, "sofia", "sofia maya", "Performance engineering", "Medium", [
        ("Interpretation", "Budgets are planning allocations, not observed end-to-end measurements. The p95 release target is strictly below 2500 milliseconds. Sofia Ortiz owns performance verification."),
    ], ["Atlas p95 latency target below 2500 ms."], table=[["Component", "Budget ms", "Evidence status"], ["Authentication", 100, "Planning allocation"], ["Retrieval", 350, "Planning allocation"], ["Model generation", 1800, "Planning allocation"], ["Rendering and transport", 200, "Planning allocation"], ["Total allocation", "=SUM(B2:B5)", "Calculated planning budget; not measured latency"], ["Release p95 threshold", 2500, "Observed p95 must be strictly below threshold"]]),
    spec("beacon_commercial_terms.docx", "Beacon Route — unsigned three-depot estimate", "Beacon Route", 4, "emma", "emma priya", "Commercial proposal", "Medium", [
        ("Draft status", "UNSIGNED ESTIMATE. This September 4 proposal assumes North, Central, and South depots. It has not been accepted by Beacon Route and is not a purchase commitment."),
        ("Proposed fee", "The estimated fixed fee for three depots is USD 62,000 excluding tax. Emma Laurent owns commercial negotiation. Any reduction to two depots requires a revised scope and price."),
        ("Exclusions", "Automatic vehicle dispatch, weather vendor subscription charges, and weekend incident support are excluded from this estimate. These exclusions are not evidence that separate vendor retention terms have been agreed."),
    ], ["Unsigned September 4 estimate: USD 62,000 for three depots, later superseded."]),
    spec("beacon_pilot_review.ppt", "Beacon Route — forecast evaluation review", "Beacon Route", 11, "noah", "noah priya jules", "Model evaluation", "High", [
        ("Baseline and holdout", "The seasonal spreadsheet baseline had 16 percent mean absolute percentage error. The deduplicated model achieved 11.4 percent on eight frozen holdout weeks across North and Central depots. No live labor-cost savings have been measured."),
        ("Acceptance", "The agreed pilot threshold is at most 12 percent forecast error. Forecasts must arrive by 06:00 local depot time. Missing-input runs must show a warning and fall back to the seasonal baseline."),
        ("Human control", "Jules Park leads planner rehearsals. A human planner approves every forecast before staffing changes. Live vehicle dispatch remains outside this pilot."),
    ], ["Beacon evaluated error 11.4%, historical baseline 16%, threshold 12%."], notes="Noah Brooks owns error measurement. The observed result is offline; commercial and vendor approvals are separate gates."),
    spec("beacon_quality_register.xls", "Beacon Route — data quality register", "Beacon Route", 12, "noah", "noah jules priya", "Data quality", "High", [
        ("Register interpretation", "This register distinguishes measured sample results from acceptance thresholds. Noah Brooks owns the connector and data-quality evidence. Jules Park owns closure calendars."),
    ], ["Central duplicate shipment IDs dropped from 37 to 0 in verified sample."], table=[["Check", "Observed result", "Acceptance rule", "Owner"], ["Central duplicate shipment IDs", "0 after cleanup; 37 before", "0 duplicates using shipment ID plus depot", "Noah Brooks"], ["History coverage", "12 complete weeks", "At least 12 complete weeks", "Noah Brooks"], ["Central holiday annotations", "2 unconfirmed closures", "Jules validates missing annotations before pilot", "Jules Park"], ["Weather vendor deletion SLA", "Unknown; vendor response pending", "Signed terms required before ongoing feed", "Emma Laurent"]]),
    spec("beacon_dispatch_handoff.xlsx", "Beacon Route — planner handoff checklist", "Beacon Route", 15, "jules", "jules priya noah", "Operations", "Medium", [
        ("Handoff boundaries", "This checklist supports manual staffing planning at North and Central depots. It does not authorize automatic vehicle dispatch. Jules Park coordinates morning rehearsals."),
    ], ["Human planners approve Beacon forecasts; missing inputs use warned baseline fallback.", "Approved handoff error threshold is 10 percent, conflicting with unreconciled 12-percent decision."], table=[["Step", "Due or trigger", "Required action", "Owner"], ["Forecast delivery", "06:00 local depot time", "Publish forecast and input freshness", "Noah Brooks"], ["Missing input", "Any required feed absent", "Warn planner and use seasonal baseline", "Noah Brooks"], ["Planner decision", "Before staffing changes", "Approve or reject forecast manually", "Jules Park"], ["Holdout acceptance", "Approved handoff requirement", "Error at most 10 percent; differs from September 11 approval", "Priya Raman"], ["Vendor feed activation", "After signed terms", "Obtain retention agreement", "Emma Laurent"]]),
    spec("cedar_privacy_review.doc", "Cedar Vale — draft privacy review", "Cedar Vale", 5, "nina", "nina ines oscar", "Privacy", "High", [
        ("Draft proposal", "DRAFT FOR REVIEW. The proposed raw-prompt retention period is seven days. This draft is not an approved policy. Nina Shah will decide after the privacy review."),
        ("Data limits", "The policy-search corpus excludes patient records. Input prompts may accidentally contain identifiers, so redaction must occur before model access. Do not use raw prompts to populate analytics dashboards."),
        ("Unresolved policy", "The approved handbook extracts do not define paid parental leave eligibility for temporary contractors. Nina Shah owns that policy clarification; retrieval cannot manufacture an eligibility rule."),
    ], ["Seven-day raw-prompt retention is an unapproved draft, later superseded."]),
    spec("cedar_access_design.pptx", "Cedar Vale — policy access design", "Cedar Vale", 8, "oscar", "oscar elliot nina", "Identity architecture", "Critical", [
        ("Claims and source labels", "Every policy chunk carries clinic_id and policy_group. The authenticated identity token supplies the user's allowed values. Apply both filters before retrieval ranking and before model context assembly."),
        ("Fail closed", "Missing identity claims, expired credentials, and a revoked policy group deny access. Never retry against a global collection. Oscar Bell owns access middleware and regressions."),
        ("Policy approval boundary", "Nina Shah approves published policy content. Corrections and unanswered questions enter a human review queue. They must not silently rewrite approved source documents."),
    ], ["Cedar filters both clinic_id and policy_group before model context."]),
    spec("cedar_operations.ppt", "Cedar Vale — incident containment runbook", "Cedar Vale", 16, "elliot", "elliot oscar ines theo", "Incident response", "Critical", [
        ("P1 containment", "For suspected cross-clinic disclosure, disable the assistant endpoint, invalidate the role cache, and preserve redacted request IDs and policy-version evidence. Oscar Bell owns containment."),
        ("Communication and recovery", "Ines Duarte owns client communication. Elliot Reed owns triage. Restore only after the revoked-group regression passes and Theo Haddad completes an independent security review."),
        ("Coverage limitation", "No signed weekend support SLA or staffed on-call rota is documented. The nine-minute endpoint shutdown in the tabletop is an observation, not a guaranteed response time. Route contractual coverage questions to Ines Duarte."),
    ], ["Cedar recovery requires revoked-group regression and Theo Haddad security review."], notes="This is a tabletop runbook. Never describe the drill as a real patient-data incident."),
    spec("cedar_eval_scores.xls", "Cedar Vale — frozen policy evaluation scores", "Cedar Vale", 13, "elliot", "elliot nina oscar", "Answer quality", "High", [
        ("Evaluation protocol", "Elliot Reed evaluated a frozen set of 80 questions against approved policy versions. These are offline evaluation results, not a guarantee about production behavior."),
    ], ["Cedar eval: 80 questions; 0/8 cross-clinic disclosures; 58/60 supported answer cases; 12/12 gaps routed."], table=[["Category", "Cases", "Observed outcome", "Acceptance threshold"], ["Answerable policy questions", 60, "58 supported; 2 need review", "At least 95 percent supported claims"], ["Deliberately missing policies", 12, "12 routed for clarification", "All known gaps routed"], ["Cross-clinic access attempts", 8, "0 disclosures", "Zero disclosures"], ["Total question count", "=SUM(B2:B4)", "Frozen before evaluation", "80 questions"]]),
]


EVALUATION = [
    {"id": "atlas-current-date", "question": "What is Atlas Forge's current launch date, and did it change from the kickoff assumption?", "expected_status": "answered", "expected_sources": ["atlas_03_security_gate.md"], "required_terms": ["November", "October"], "forbidden_terms": []},
    {"id": "atlas-logging", "question": "How long can Atlas keep diagnostic logs, where are they kept, and may they contain raw prompts?", "expected_status": "answered", "expected_sources": ["atlas_security_addendum.docx"], "required_terms": ["30", "Frankfurt"], "forbidden_terms": []},
    {"id": "atlas-mixed-release", "question": "Explain Atlas's staged rollout and the latency stop condition, including whether the workbook contains measured latency.", "expected_status": "answered", "expected_sources": ["atlas_release_plan.pptx", "atlas_latency_budget.xlsx"], "required_terms": ["5", "25"], "forbidden_terms": []},
    {"id": "atlas-window-gap", "question": "What exact rollback maintenance window has the Atlas factory approved?", "expected_status": "needs_routing", "expected_sources": ["atlas_04_delivery.md"], "required_terms": [], "forbidden_terms": ["02:00-04:00"]},
    {"id": "beacon-price-change", "question": "Which Beacon pilot price is signed, and why does an older proposal have a different figure?", "expected_status": "answered", "expected_sources": ["beacon_04_commercial.md", "beacon_commercial_terms.docx"], "required_terms": ["48", "62"], "forbidden_terms": []},
    {"id": "beacon-paraphrase", "question": "If a required feed does not arrive before Beacon's morning planning run, what should the planner receive?", "expected_status": "answered", "expected_sources": ["beacon_dispatch_handoff.xlsx"], "required_terms": ["baseline", "warning"], "forbidden_terms": []},
    {"id": "beacon-vendor-gap", "question": "What is the weather vendor's signed deletion SLA for Beacon Route?", "expected_status": "needs_routing", "expected_sources": ["beacon_02_data_review.md", "beacon_quality_register.xls"], "required_terms": [], "forbidden_terms": ["72 hours"]},
    {"id": "cedar-retention-change", "question": "Is Cedar's seven-day raw-prompt retention proposal still valid? State the approved rule and its owner.", "expected_status": "answered", "expected_sources": ["cedar_03_validation.md", "cedar_privacy_review.doc"], "required_terms": ["24", "Nina"], "forbidden_terms": []},
    {"id": "cedar-controls", "question": "How does Cedar prevent staff from seeing another clinic's policies, and what happens when identity claims are absent?", "expected_status": "answered", "expected_sources": ["cedar_access_design.pptx"], "required_terms": ["clinic", "policy"], "forbidden_terms": []},
    {"id": "beacon-conflicting-threshold", "question": "What is Beacon Route's approved forecast error threshold? Explain the conflicting acceptance records and who must resolve them.", "expected_status": "partial", "expected_sources": ["beacon_03_pilot_acceptance.md", "beacon_dispatch_handoff.xlsx"], "required_terms": ["10", "12", "Priya"], "forbidden_terms": []},
    {"id": "unrelated-abstention", "question": "What was the winning time in the 1978 lunar yacht championship?", "expected_status": "needs_routing", "expected_sources": [], "required_terms": [], "forbidden_terms": []},
    {"id": "cedar-eval-results", "question": "What were Cedar's offline policy evaluation results for missing policies and cross-clinic access attempts?", "expected_status": "answered", "expected_sources": ["cedar_eval_scores.xls"], "required_terms": ["12", "8"], "forbidden_terms": []},
]

# Daily canary uses the first three cases: an answer, an unresolved contradiction,
# and a knowledge gap with an attributable owner. The full suite stays held out.
EVALUATION = [EVALUATION[0], EVALUATION[9], EVALUATION[6]] + [case for index, case in enumerate(EVALUATION) if index not in {0, 9, 6}]


def headers(item: dict) -> list[str]:
    return [f"Title: {item['title']}", f"Date: {item['date']}", f"Author: {item['author']}", f"Attendees: {'; '.join(item['attendees'])}", f"Client: {item['client']}", f"Domain: {item['domain']}", f"Priority: {item['priority']}", "Classification: SYNTHETIC — fictional evaluation material"]


def write_markdown(path: Path, item: dict) -> None:
    body = [f"# {item['title']}", "", "\n".join(headers(item)[1:]), "", "Transcript — edited for clarity; all people and engagements are fictional.", ""]
    for speaker, words in item["blocks"]:
        body += [f"{speaker}: {words}", ""]
    # Hash the same source bytes on Windows and after Git's LF checkout.
    path.write_text("\n".join(body), encoding="utf-8", newline="\n")


def write_word(path: Path, item: dict) -> None:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(7)
    # Keep reusable Office defaults from introducing a decorative title rule.
    title_style = document.styles["Title"]
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.font.underline = False
    title_style.font.size = Pt(23)
    for element in document.styles.element.iter(qn("w:pBdr")):
        element.getparent().remove(element)
    visible_title = item["title"].replace(" — ", " ").replace("-", " ")
    document.add_heading(visible_title, 0)
    for line in headers(item)[1:]:
        p = document.add_paragraph(line)
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string("475569")
    for title, body in item["blocks"]:
        document.add_heading(title, 1)
        document.add_paragraph(body)
    props = document.core_properties
    props.title, props.author, props.subject = item["title"], item["author"], item["domain"]
    props.created = props.modified = datetime.fromisoformat(item["date"] + "T12:00:00")
    document.save(path)


def write_slides(path: Path, item: dict) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    pages = [(item["title"], "\n".join(headers(item)[1:]))] + item["blocks"]
    for index, (title, body) in enumerate(pages):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F8FAFC")
        heading = slide.shapes.add_textbox(Inches(.7), Inches(.65), Inches(12), Inches(1.3)).text_frame
        heading.word_wrap = True
        heading.text = title
        heading.paragraphs[0].font.size = Pt(30 if index else 28)
        heading.paragraphs[0].font.bold = True
        heading.paragraphs[0].font.color.rgb = RGBColor.from_string("0F172A")
        content = slide.shapes.add_textbox(Inches(.75), Inches(2.1), Inches(11.8), Inches(4.4)).text_frame
        content.word_wrap = True
        content.text = body
        for paragraph in content.paragraphs:
            paragraph.font.size = Pt(18 if index == 0 else 25)
            paragraph.font.color.rgb = RGBColor.from_string("334155")
            paragraph.space_after = Pt(10)
        footer = slide.shapes.add_textbox(Inches(.75), Inches(7), Inches(11), Inches(.3)).text_frame
        footer.text = f"RELAY AI  /  {item['client'].upper()}  /  SYNTHETIC DATA  /  {index + 1}"
        footer.paragraphs[0].font.size = Pt(10)
        if index == len(pages) - 1 and item["notes"]:
            slide.notes_slide.notes_text_frame.text = item["notes"]
    props = presentation.core_properties
    props.title, props.author, props.subject = item["title"], item["author"], item["domain"]
    props.created = props.modified = datetime.fromisoformat(item["date"] + "T12:00:00")
    presentation.save(path)


def write_sheet(path: Path, item: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    about = workbook.active
    about.title = "Read me"
    for row in headers(item):
        about.append([row])
    about.append([])
    for title, body in item["blocks"]:
        about.append([f"{title}: {body}"])
    about.column_dimensions["A"].width = 100
    for row in about:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name="Calibri", size=11)
        about.row_dimensions[row[0].row].height = 45 if row[0].row > 8 else 28
    about["A1"].font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    sheet = workbook.create_sheet("Register")
    for row in item["table"]:
        sheet.append(row)
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Calibri", size=11, bold=cell.row == 1, color="FFFFFF" if cell.row == 1 else "1E293B")
            cell.fill = PatternFill("solid", fgColor="0F766E" if cell.row == 1 else ("F1F5F9" if cell.row % 2 else "FFFFFF"))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row[0].row].height = 45
    for column in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 37 if column != 1 else 34
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for tab in workbook:
        tab.sheet_properties.pageSetUpPr.fitToPage = True
        tab.page_setup.orientation = "landscape"
        tab.page_setup.paperSize = tab.PAPERSIZE_A4
        tab.page_setup.fitToWidth = 1
        tab.page_setup.fitToHeight = 1
        tab.print_options.horizontalCentered = True
        tab.print_area = tab.dimensions
    props = workbook.properties
    props.title, props.creator, props.subject = item["title"], item["author"], item["domain"]
    props.created = props.modified = datetime.fromisoformat(item["date"] + "T12:00:00")
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--soffice", help="Path to LibreOffice soffice executable")
    parser.add_argument("--output", type=Path, default=ROOT / "data/corpus")
    args = parser.parse_args()
    executable = discover_soffice(args.soffice)
    args.output.mkdir(parents=True, exist_ok=True)
    manifest: list[dict] = []
    writers = {".md": write_markdown, ".docx": write_word, ".pptx": write_slides, ".xlsx": write_sheet}
    legacy = {".doc": (".docx", "doc:MS Word 97"), ".ppt": (".pptx", "ppt:MS PowerPoint 97"), ".xls": (".xlsx", "xls:MS Excel 97")}
    for item in SPECS:
        path = args.output / item["filename"]
        if path.suffix in legacy:
            modern_extension, filter_name = legacy[path.suffix]
            with tempfile.TemporaryDirectory(prefix="relay-generate-") as temporary:
                temporary_path = Path(temporary)
                modern = temporary_path / (path.stem + modern_extension)
                writers[modern_extension](modern, item)
                profile = (temporary_path / "profile").as_uri()
                converted_directory = temporary_path / "legacy"
                converted_directory.mkdir()
                converted = converted_directory / path.name
                result = subprocess.run([executable, f"-env:UserInstallation={profile}", "--headless", "--convert-to", filter_name, "--outdir", str(converted_directory), str(modern)], capture_output=True, text=True, timeout=90)
                if result.returncode or not converted.exists() or converted.read_bytes()[:8] != bytes.fromhex("D0CF11E0A1B11AE1"):
                    raise RuntimeError(f"LibreOffice failed to create genuine {path.suffix} file: {path.name}; {result.stdout} {result.stderr}")
                path.write_bytes(converted.read_bytes())
        else:
            writers[path.suffix](path, item)
        parsed = extract_document(path, soffice_path=executable)
        assert parsed.author == item["author"], f"Author changed during round-trip: {path.name}"
        assert parsed.date == item["date"], f"Source date changed during round-trip: {path.name}"
        assert parsed.attendees == item["attendees"], f"Attendees changed during round-trip: {path.name}"
        manifest.append({key: value for key, value in item.items() if key not in {"blocks", "table", "notes"}} | {"sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "chunk_count": len(parsed.chunks), "warnings": parsed.warnings})
        print(f"Created {path.name}: {len(parsed.chunks)} source chunks")
    expected = {".md": 12, ".doc": 2, ".docx": 2, ".ppt": 2, ".pptx": 2, ".xls": 2, ".xlsx": 2}
    actual = Counter(Path(item["filename"]).suffix for item in SPECS)
    assert actual == expected, actual
    (args.output.parent / "manifest.json").write_text(json.dumps({"description": "Fictional Relay AI consulting corpus. No real client information.", "generator": "scripts/generate_corpus.py", "files": manifest}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    (args.output.parent / "evaluation.json").write_text(json.dumps(EVALUATION, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print("Verified 24 documents with source attribution; wrote manifest and held-out evaluation outside corpus.")


if __name__ == "__main__":
    main()
