from __future__ import annotations

import json
import hashlib
from pathlib import Path
import subprocess

import pytest

from backend.app.extractors import ExtractionError, discover_soffice, extract_document


ROOT = Path(__file__).resolve().parents[1]


def test_generated_markdown_has_canonical_lf_bytes(tmp_path):
    from scripts.generate_corpus import SPECS, write_markdown

    path = tmp_path / "meeting.md"
    write_markdown(path, SPECS[0])
    content = path.read_bytes()
    assert b"\n" in content
    assert b"\r" not in content
    assert content == (ROOT / "data/corpus" / SPECS[0]["filename"]).read_bytes()


def test_markdown_preserves_identity_and_line_locations(tmp_path):
    path = tmp_path / "meeting.md"
    path.write_text("# Release review\n\nAuthor: Asha Patel\nDate: 2026-09-10\nAttendees: Asha Patel; Morgan Lee\n\nAsha Patel: We launch after approval.\n", encoding="utf-8")
    result = extract_document(path)
    assert (result.title, result.author, result.date) == ("Release review", "Asha Patel", "2026-09-10")
    assert result.attendees == ["Asha Patel", "Morgan Lee"]
    assert result.chunks[-1].locator == "Lines 7-7"
    assert result.chunks[-1].text == "Asha Patel: We launch after approval."
    assert not result.warnings


def test_missing_author_does_not_invent_identity_or_use_default(tmp_path):
    from docx import Document

    path = tmp_path / "unattributed.docx"
    document = Document()
    document.add_paragraph("An unattributed procedure.")
    document.save(path)
    result = extract_document(path)
    assert result.author is None
    assert result.date is None
    assert any("author is missing" in message for message in result.warnings)
    assert any("Source date is missing" in message for message in result.warnings)


def test_word_keeps_body_and_table_order_and_visible_author_wins(tmp_path):
    from docx import Document

    path = tmp_path / "procedure.docx"
    document = Document()
    document.core_properties.author = "Template Author"
    document.add_paragraph("Author: Asha Patel\nDate: 2026-09-10")
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Condition", "Action"
    table.cell(1, 0).text, table.cell(1, 1).text = "Missing claim", "Deny access"
    document.add_paragraph("After table")
    document.save(path)
    result = extract_document(path)
    assert result.author == "Asha Patel"
    assert result.chunks[1].text == "Before table"
    assert result.chunks[3].locator == "Table 1, row 2, cells 1-2"
    assert "Deny access" in result.chunks[3].text
    assert result.chunks[4].text == "After table"
    assert all("page" not in chunk.locator.lower() for chunk in result.chunks)


def test_slides_include_group_text_table_cells_and_speaker_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "review.pptx"
    presentation = Presentation()
    presentation.core_properties.author = "Asha Patel"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text = "Date: 2026-09-10\nGrouped decision"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Owner", "Morgan Lee"
    slide.notes_slide.notes_text_frame.text = "Approval remains pending."
    presentation.save(path)
    result = extract_document(path)
    assert result.date == "2026-09-10"
    assert any("group" in chunk.locator and "Grouped decision" in chunk.text for chunk in result.chunks)
    assert any("table row 1" in chunk.locator and "Morgan Lee" in chunk.text for chunk in result.chunks)
    assert any(chunk.locator == "Slide 1, speaker notes" and "pending" in chunk.text for chunk in result.chunks)


def test_spreadsheet_formula_without_cached_value_is_explicit(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "budget.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Costs"
    sheet.append(["Author: Asha Patel"])
    sheet.append(["Date: 2026-09-10"])
    sheet.append(["Estimate", 12])
    sheet.append(["Total", "=SUM(B3:B3)"])
    workbook.save(path)
    result = extract_document(path)
    assert result.author == "Asha Patel"
    assert result.date == "2026-09-10"
    assert result.chunks[-1].locator == "Sheet 'Costs', cells A4:B4"
    assert "=SUM(B3:B3) [cached result unavailable]" in result.chunks[-1].text
    assert any("Costs!B4" in message and "not evaluated" in message for message in result.warnings)


@pytest.mark.parametrize("name,content,fragment", [("wrong.pdf", b"not pdf", "Unsupported"), ("empty.md", b"", "No extractable"), ("broken.docx", b"not an OOXML file", "Cannot extract")])
def test_invalid_sources_are_actionable(tmp_path, name, content, fragment):
    path = tmp_path / name
    path.write_bytes(content)
    with pytest.raises(ExtractionError, match=fragment):
        extract_document(path)


def test_invalid_explicit_soffice_does_not_silently_choose_another(tmp_path):
    with pytest.raises(ExtractionError, match="SOFFICE_PATH"):
        discover_soffice(str(tmp_path / "absent-soffice"))


def test_legacy_conversion_timeout_is_reported(tmp_path, monkeypatch):
    path = tmp_path / "source.doc"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1"))
    monkeypatch.setattr("backend.app.extractors.discover_soffice", lambda configured: "soffice")

    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("soffice", 90)

    monkeypatch.setattr("backend.app.extractors.subprocess.run", timeout)
    with pytest.raises(ExtractionError, match="timed out after 90 seconds"):
        extract_document(path)


def test_legacy_conversion_requires_output_file(tmp_path, monkeypatch):
    path = tmp_path / "source.doc"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1"))
    monkeypatch.setattr("backend.app.extractors.discover_soffice", lambda configured: "soffice")
    monkeypatch.setattr("backend.app.extractors.subprocess.run", lambda *args, **kwargs: subprocess.CompletedProcess(args[0], 0, "", ""))
    with pytest.raises(ExtractionError, match="could not convert"):
        extract_document(path)


def test_corpus_contains_all_required_distinct_formats():
    from collections import Counter

    manifest = json.loads((ROOT / "data/manifest.json").read_text(encoding="utf-8"))
    assert len(manifest["files"]) == 24
    assert Counter(Path(item["filename"]).suffix for item in manifest["files"]) == {".md": 12, ".doc": 2, ".docx": 2, ".ppt": 2, ".pptx": 2, ".xls": 2, ".xlsx": 2}
    assert len({item["sha256"] for item in manifest["files"]}) == 24
    for item in manifest["files"]:
        path = ROOT / "data/corpus" / item["filename"]
        assert path.is_file()
        assert hashlib.sha256(path.read_bytes()).hexdigest() == item["sha256"], f"Manifest hash is stale: {path.name}"
        if path.suffix == ".md":
            assert b"\r" not in path.read_bytes(), f"Transcript must use canonical LF bytes: {path.name}"
        if path.suffix in {".doc", ".ppt", ".xls"}:
            assert path.read_bytes()[:8] == bytes.fromhex("D0CF11E0A1B11AE1"), f"{path.name} is not an Office compound binary"


@pytest.mark.parametrize("filename", sorted(path.name for path in (ROOT / "data/corpus").iterdir()))
def test_every_committed_source_round_trips_attribution(filename):
    manifest_path = ROOT / "data/manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    try:
        soffice = discover_soffice()
    except ExtractionError:
        soffice = None
    item = next(item for item in manifest["files"] if item["filename"] == filename)
    path = ROOT / "data/corpus" / filename
    if path.suffix in {".doc", ".ppt", ".xls"} and not soffice:
        pytest.skip("LibreOffice is unavailable; real binary header is checked by corpus manifest test.")
    result = extract_document(path, soffice_path=soffice)
    assert result.author == item["author"], path.name
    assert result.date == item["date"], path.name
    assert result.attendees == item["attendees"], path.name
    assert result.chunks
    assert all(chunk.locator and chunk.text for chunk in result.chunks)
