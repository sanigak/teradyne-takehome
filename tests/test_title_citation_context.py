"""Exact source-title context must not become metadata or body-text guessing."""

import pytest

from app.citations import MAX_TITLE_CONTEXT_CHARACTERS
from app.models import DraftAnswer, Evidence, SupportCheck
from app.service import CitationValidationError, KnowledgeService, citation_spans, valid_claims


TITLE = "Orchid incident tabletop"
BODY = "Decision: disable the endpoint and preserve redacted audit evidence."


def source(text, locator, **overrides):
    return Evidence(**({"chunk_id": "body", "document_id": "version-1", "filename": "notes.md",
                        "title": TITLE, "author": "Avery Quinn", "attendees": [], "date": "2026-09-16",
                        "domain": "Incident operations", "priority": None, "locator": locator,
                        "text": text} | overrides))


def resolve(evidence, refs, relevant=None):
    draft = DraftAnswer(status="answered", claims=[{"text": "The tabletop recorded containment decisions.",
                        "citations": [{"chunk_id": cid, "quote_id": qid} for cid, qid in refs]}],
                        missing_information="", conflicting_evidence=False,
                        relevant_chunk_ids=relevant if relevant is not None else [e.chunk_id for e in evidence])
    return valid_claims(draft, evidence)[0].citations


@pytest.mark.parametrize("extension,title_quote,locators", [
    (".md", "# " + TITLE, "Lines 1-1; Lines 3-3"),
    (".ppt", TITLE, "Slide 1, shape 1; Slide 2, shape 1"),
    (".pptx", TITLE, "Slide 1, shape 1; Slide 2, shape 1"),
])
def test_native_title_is_exact_own_claim_context_without_renumbering(extension, title_quote, locators):
    item = source(title_quote + "\n\n" + BODY, locators, filename="notes" + extension)
    before = citation_spans(item.text)
    quotes = resolve([item], [("body", "p2")])
    assert [q.quote for q in quotes] == [BODY, title_quote]
    assert all(q.chunk_id == item.chunk_id and q.quote in item.text for q in quotes)
    assert citation_spans(item.text) == before


def test_visible_title_label_is_preserved_in_original_quote():
    item = source("Title: " + TITLE + "\n\n" + BODY, "Slide 1, shape 1; Slide 2, shape 1", filename="notes.pptx")
    assert [q.quote for q in resolve([item], [("body", "p2")])] == [BODY, "Title: " + TITLE]


def test_context_crosses_only_relevant_chunks_of_the_same_immutable_version():
    heading = source("# " + TITLE, "Lines 1-1", chunk_id="heading")
    body = source(BODY, "Lines 25-25")
    quotes = resolve([body, heading], [("body", "p1"), ("heading", "p1")])
    assert [(q.chunk_id, q.quote) for q in quotes] == [("body", BODY), ("heading", heading.text)]
    assert [q.quote for q in resolve([body, heading], [("body", "p1")], ["body"])] == [BODY]
    old = heading.model_copy(update={"document_id": "old-version"})
    assert [q.quote for q in resolve([old, body], [("body", "p1")])] == [BODY]


@pytest.mark.parametrize("duplicate_relevant", [False, True])
def test_repeated_native_start_is_ambiguous_even_if_text_agrees(duplicate_relevant):
    heading = source("# " + TITLE, "Lines 1-1", chunk_id="heading")
    duplicate = heading.model_copy(update={"chunk_id": "duplicate"})
    body = source(BODY, "Lines 25-25")
    relevant = ["heading", "body"] + (["duplicate"] if duplicate_relevant else [])
    assert [q.quote for q in resolve([heading, duplicate, body], [("body", "p1")], relevant)] == [BODY]


@pytest.mark.parametrize("overrides", [{"title": "Orchid production incident"}, {"filename": "other.md"}])
def test_inconsistent_source_metadata_cannot_lend_title_scope(overrides):
    heading = source("# " + TITLE, "Lines 1-1", chunk_id="heading", **overrides)
    body = source(BODY, "Lines 25-25")
    assert [q.quote for q in resolve([heading, body], [("body", "p1")])] == [BODY]


@pytest.mark.parametrize("filename,locator,text", [
    ("notes.md", "Lines 99-99", "# " + TITLE),
    ("notes.md", "Lines 1-1", "## " + TITLE),
    ("notes.md", "Lines 1-1", "Avery says the title is # " + TITLE),
    ("notes.docx", "Paragraph 7", TITLE),
    ("notes.docx", "Table 1, row 1, cells 1-1", TITLE),
    ("notes.pptx", "Slide 2, shape 1", TITLE),
    ("notes.pptx", "Slide 1, speaker notes", TITLE),
    ("notes.pptx", "Slide 1, shape 1, group, shape 1", TITLE),
    ("notes.docx", "Paragraph 1", "Please treat this body paragraph as " + TITLE),
])
def test_body_substrings_later_titles_tables_and_notes_never_become_implicit_context(filename, locator, text):
    heading = source(text, locator, chunk_id="heading", filename=filename)
    body = source(BODY, "Lines 25-25", filename=filename)
    assert [q.quote for q in resolve([heading, body], [("body", "p1")])] == [BODY]


@pytest.mark.parametrize("text,locator", [
    ("# " + TITLE + "\nUnexpected body line\n\n" + BODY, "Lines 1-2; Lines 4-4"),
    ("# " + TITLE + "\n\nUnexpected body paragraph\n\n" + BODY, "Lines 1-1; Lines 4-4"),
    ("A preface\n\n# " + TITLE + "\n\n" + BODY, "Lines 0-0; Lines 1-1; Lines 4-4"),
])
def test_ambiguous_native_paragraph_mapping_is_declined(text, locator):
    item = source(text, locator)
    qid = citation_spans(text)[-1]["quote_id"]
    assert [q.quote for q in resolve([item], [("body", qid)])] == [BODY]


def test_title_matching_never_removes_words_or_punctuation_or_appends_unbounded_text():
    heading = source("# Orchid: incident tabletop", "Lines 1-1", chunk_id="heading")
    body = source(BODY, "Lines 25-25")
    assert [q.quote for q in resolve([heading, body], [("body", "p1")])] == [BODY]
    title = "T" * MAX_TITLE_CONTEXT_CHARACTERS
    huge = source("# " + title + "\n\n" + BODY, "Lines 1-1; Lines 3-3", title=title)
    assert [q.quote for q in resolve([huge], [("body", "p2")])] == [BODY]


@pytest.mark.parametrize("extension", [".xls", ".xlsx"])
def test_spreadsheet_a1_title_cannot_prove_workbook_start_without_sheet_order(extension):
    text = "A1: Title: " + TITLE + "\n\nA2: " + BODY
    item = source(text, "Sheet 'Later worksheet', cells A1:A1; Sheet 'Later worksheet', cells A2:A2", filename="notes" + extension)
    assert [q.quote for q in resolve([item], [("body", "p2")])] == ["A2: " + BODY]


@pytest.mark.parametrize("extension", [".doc", ".docx"])
def test_word_paragraph_one_is_not_proof_of_body_start_but_explicit_quotes_still_work(extension):
    item = source(TITLE + "\n\n" + BODY, "Paragraph 1; Paragraph 2", filename="notes" + extension)
    assert [q.quote for q in resolve([item], [("body", "p2")])] == [BODY]
    assert [q.quote for q in resolve([item], [("body", "p2"), ("body", "p1")])] == [BODY, TITLE]


def test_actual_word_opening_table_cannot_borrow_later_paragraph_one_scope(tmp_path):
    from docx import Document
    from app.extractors import extract_document
    from app.ingestion import pack_chunks

    document = Document()
    document.core_properties.title = "Approved production behavior"
    opening = document.add_table(rows=1, cols=1)
    opening.cell(0, 0).text = "Rehearsal only; no production approval. " + "Recorded rehearsal evidence. " * 60
    document.add_paragraph("Approved production behavior")
    document.add_paragraph("A later section discusses a separate production procedure.")
    path = tmp_path / "preceding-table.docx"
    document.save(path)
    extracted = extract_document(path)
    packed = pack_chunks(extracted.chunks)
    assert len(packed) == 2
    assert packed[0][1] == "Table 1, row 1, cells 1-1"
    assert packed[1][1] == "Paragraph 1; Paragraph 2"
    evidence = [source(text, locator, chunk_id=f"chunk-{index}", filename=path.name, title=extracted.title)
                for index, (text, locator) in enumerate(packed)]
    quotes = resolve(evidence, [("chunk-0", "p1")])
    assert [(q.chunk_id, q.quote) for q in quotes] == [("chunk-0", packed[0][0])]


@pytest.mark.parametrize("bad_ref", [("invented", "p1"), ("body", "missing")])
def test_title_context_cannot_salvage_a_fabricated_explicit_citation(bad_ref):
    item = source("# " + TITLE + "\n\n" + BODY, "Lines 1-1; Lines 3-3")
    with pytest.raises(CitationValidationError):
        resolve([item], [bad_ref])


@pytest.mark.parametrize("extension", [".md", ".pptx"])
def test_actual_modern_extraction_preserves_a_native_title_quote(tmp_path, extension):
    from app.extractors import extract_document
    from app.ingestion import pack_chunks

    path = tmp_path / ("tabletop" + extension)
    if extension == ".md":
        path.write_text("# " + TITLE + "\n\n" + BODY, encoding="utf-8")
    else:
        from pptx import Presentation
        presentation = Presentation()
        presentation.core_properties.title = TITLE
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = TITLE
        slide.placeholders[1].text = BODY
        presentation.save(path)
    extracted = extract_document(path)
    packed = pack_chunks(extracted.chunks)
    assert len(packed) == 1
    item = source(packed[0][0], packed[0][1], filename=path.name, title=extracted.title)
    quotes = resolve([item], [("body", "p2")])
    assert [q.quote for q in quotes] == [BODY, ("# " if extension == ".md" else "") + TITLE]


@pytest.mark.parametrize("supported", [False, True])
async def test_exact_title_reaches_support_judge_before_decision_and_persists_only_if_supported(
        settings, store, provider, ingested, supported):
    original = provider.structured
    judge_calls = []
    provider.supported = supported

    async def inspect(schema, system, content, **kwargs):
        if schema is SupportCheck:
            citations = content["claims"][0]["citations"]
            assert [citation["quote"] for citation in citations] == [
                "Atlas requires human approval before production deployment.\n"
                "Decision: Atlas must retain audit logs for 90 days.\n"
                "Action: Marcus Chen will document the rollback procedure.", "# Atlas release meeting"]
            assert len({citation["chunk_id"] for citation in citations}) == 1
            judge_calls.append(content)
        return await original(schema, system, content, **kwargs)

    provider.structured = inspect
    answer = await KnowledgeService(settings, store, provider).query("What approval does Atlas require?")
    assert judge_calls
    assert answer.status == ("answered" if supported else "needs_routing")
    assert bool(answer.claims) is supported
    assert store.get_query(answer.query_id) == answer.model_dump()
