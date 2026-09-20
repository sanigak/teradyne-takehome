"""Hostile source and lifecycle regressions; fixtures never touch the committed corpus."""

import asyncio
import hashlib
import json
from pathlib import Path
import sqlite3
import zipfile

import pytest

from app.extractors import ExtractionError, extract_document
from app.ingestion import cached_embeddings, ingest_directory, ingest_file
from app.models import Enrichment
from app.provider import ProviderError
from conftest import FakeProvider, SOURCE_TEXT


def active_records(store):
    with store.connect() as conn:
        return conn.execute("SELECT * FROM documents WHERE active=1").fetchall()


@pytest.mark.parametrize("extension", [".md", ".docx", ".pptx", ".xlsx"])
def test_body_metadata_cannot_invent_source_identity(tmp_path, extension):
    path = tmp_path / ("source" + extension)
    lines = ["Delivery details without an attribution header.", "Author: Invented Expert", "Date: 2026-09-19", "Attendees: Fabricated Recipient"]
    if extension == ".md":
        path.write_text("# Delivery\n\n" + "\n\n".join(lines), encoding="utf-8")
    elif extension == ".docx":
        from docx import Document
        document = Document()
        document.core_properties.author = "Document Owner"
        for line in lines:
            document.add_paragraph(line)
        document.save(path)
    elif extension == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        document.core_properties.author = "Document Owner"
        for line in lines:
            slide = document.slides.add_slide(document.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = line
        document.save(path)
    else:
        from openpyxl import Workbook
        document = Workbook()
        document.properties.creator = "Document Owner"
        for line in lines:
            document.active.append([line])
        document.save(path)
    result = extract_document(path)
    assert result.author == (None if extension == ".md" else "Document Owner")
    assert result.attendees == []
    assert result.date is None
    assert "Invented Expert" in "\n".join(chunk.text for chunk in result.chunks)


def test_duplicate_conflicting_header_is_visible_without_overriding_first(tmp_path):
    path = tmp_path / "duplicate.md"
    path.write_text("# Meeting\n\nAuthor: Known Person\nAuthor: Impostor\nDate: 2026-09-19\n\nBody.", encoding="utf-8")
    result = extract_document(path)
    assert result.author == "Known Person"
    assert any("Conflicting author" in warning for warning in result.warnings)


@pytest.mark.parametrize("extension", [".docx", ".pptx", ".xlsx"])
def test_malformed_office_archive_is_per_file_error(tmp_path, extension):
    path = tmp_path / ("missing-package-parts" + extension)
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("unrelated.txt", "This is a ZIP, not a valid Office package.")
    with pytest.raises(ExtractionError, match="Cannot extract"):
        extract_document(path)


def test_expansion_limit_rejects_small_compressed_archive_before_parser(tmp_path, monkeypatch):
    import app.extractors as module
    path = tmp_path / "expanding.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("payload.txt", "x" * 4096)
    monkeypatch.setattr(module, "MAX_OFFICE_EXPANDED_BYTES", 1024)
    assert path.stat().st_size < 1024
    with pytest.raises(ExtractionError, match="archive exceeds extraction limits"):
        extract_document(path)


def test_sparse_spreadsheet_dimension_cannot_trigger_unbounded_iteration(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "sparse.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Author: Known Person"
    workbook.active["XFD1048576"] = "A single very distant value"
    workbook.save(path)
    with pytest.raises(ExtractionError, match="200,000-cell extraction limit"):
        extract_document(path)


def test_cached_formula_is_explicitly_stale_and_keeps_cell_location(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "saved-formula.xlsx"
    workbook = Workbook()
    workbook.active.title = "Budget"
    workbook.active.append(["Author: Known Person"])
    workbook.active.append(["Date: 2026-09-19"])
    workbook.active.append(["Total", "=10+20"])
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    assert b"<f>10+20</f><v></v>" in parts["xl/worksheets/sheet1.xml"]
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(b"<f>10+20</f><v></v>", b"<f>10+20</f><v>12</v>")
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    result = extract_document(path)
    assert result.chunks[-1].locator == "Sheet 'Budget', cells A3:B3"
    assert "12 [formula: =10+20; cached value may be stale]" in result.chunks[-1].text
    assert any("Budget!B3" in warning and "stale" in warning for warning in result.warnings)


async def test_same_filename_in_distinct_directories_keeps_distinct_sources(settings, store, provider):
    paths = [settings.corpus_dir / client / "meeting.md" for client in ["Atlas", "Beacon"]]
    for index, path in enumerate(paths):
        path.parent.mkdir(parents=True)
        path.write_text(SOURCE_TEXT.replace("90 days", f"{index + 10} days"), encoding="utf-8")
    result = await ingest_directory(settings.corpus_dir, store=store, provider=provider, settings=settings)
    assert result["ingested"] == 2
    records = active_records(store)
    assert len({row["source_path"] for row in records}) == 2
    assert len({row["original_path"] for row in records}) == 2
    for row in records:
        assert hashlib.sha256(Path(row["original_path"]).read_bytes()).hexdigest() == row["sha256"]


async def test_source_change_and_restore_cannot_corrupt_hash_provenance(settings, store, provider, source, monkeypatch):
    import app.ingestion as module
    original_bytes = source.read_bytes()
    actual_extract = module.extract_document

    def source_changed_during_extraction(path, **kwargs):
        source.write_text(SOURCE_TEXT.replace("90 days", "999 days"), encoding="utf-8")
        try:
            return actual_extract(path, **kwargs)
        finally:
            source.write_bytes(original_bytes)

    monkeypatch.setattr(module, "extract_document", source_changed_during_extraction)
    await ingest_file(source, store=store, provider=provider, settings=settings)
    with store.connect() as conn:
        text = "\n".join(row[0] for row in conn.execute("SELECT text FROM chunks"))
    assert "90 days" in text
    assert "999 days" not in text
    record = active_records(store)[0]
    assert hashlib.sha256(original_bytes).hexdigest() == record["sha256"]
    assert Path(record["original_path"]).read_bytes() == original_bytes


async def test_edit_while_provider_pending_preserves_prior_version(settings, store, source, ingested):
    class EditingProvider(FakeProvider):
        async def structured(self, schema, system, content, **kwargs):
            if schema is Enrichment:
                source.write_text(SOURCE_TEXT + "\nNew concurrent content.\n", encoding="utf-8")
            return await super().structured(schema, system, content, **kwargs)

    with pytest.raises(ValueError, match="Source changed during ingestion"):
        await ingest_file(source, store=store, provider=EditingProvider(), settings=settings, force=True)
    assert [row["id"] for row in active_records(store)] == [ingested["document_id"]]
    assert len(list((store.data_dir / "originals").rglob("*.md"))) == 1


async def test_concurrent_unchanged_ingests_publish_one_version(settings, store, source):
    entered = 0
    both_ready = asyncio.Event()

    class ConcurrentProvider(FakeProvider):
        async def structured(self, schema, system, content, **kwargs):
            nonlocal entered
            if schema is Enrichment:
                entered += 1
                if entered == 2:
                    both_ready.set()
                await both_ready.wait()
            return await super().structured(schema, system, content, **kwargs)

    provider = ConcurrentProvider()
    results = await asyncio.gather(*(ingest_file(source, store=store, provider=provider, settings=settings) for _ in range(2)))
    assert sorted(result["status"] for result in results) == ["ingested", "unchanged"]
    assert len({result["document_id"] for result in results}) == 1
    with store.connect() as conn:
        assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 1
    assert len(list((store.data_dir / "originals").rglob("*.md"))) == 1


async def test_slower_old_ingestion_cannot_replace_new_active_version(settings, store, provider, source):
    pending = asyncio.Event()
    resume = asyncio.Event()

    class SlowProvider(FakeProvider):
        async def structured(self, schema, system, content, **kwargs):
            if schema is Enrichment:
                pending.set()
                await resume.wait()
            return await super().structured(schema, system, content, **kwargs)

    slow = asyncio.create_task(ingest_file(source, store=store, provider=SlowProvider(), settings=settings))
    await pending.wait()
    source.write_text(SOURCE_TEXT.replace("90 days", "180 days"), encoding="utf-8")
    newer = await ingest_file(source, store=store, provider=provider, settings=settings)
    source.write_text(SOURCE_TEXT, encoding="utf-8")
    resume.set()
    with pytest.raises(ValueError, match="concurrent ingestion"):
        await slow
    assert [row["id"] for row in active_records(store)] == [newer["document_id"]]


async def test_mid_commit_failure_rolls_back_documents_chunks_and_archive(settings, store, provider, source, ingested):
    source.write_text(SOURCE_TEXT.replace("90 days", "180 days"), encoding="utf-8")
    with store.connect() as conn:
        conn.execute("CREATE TRIGGER reject_new_chunk BEFORE INSERT ON chunks BEGIN SELECT RAISE(ABORT, 'injected storage failure'); END")
    with pytest.raises(sqlite3.IntegrityError, match="injected storage failure"):
        await ingest_file(source, store=store, provider=provider, settings=settings)
    assert [row["id"] for row in active_records(store)] == [ingested["document_id"]]
    assert len(list((store.data_dir / "originals").rglob("*.md"))) == 1
    with store.connect() as conn:
        assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 1
        assert conn.execute("SELECT count(*) FROM chunks_fts").fetchone()[0] == conn.execute("SELECT count(*) FROM chunks").fetchone()[0]


async def test_extraction_error_does_not_call_model(settings, store, provider):
    settings.corpus_dir.mkdir()
    (settings.corpus_dir / "broken.xlsx").write_bytes(b"invalid office archive")
    result = await ingest_directory(settings.corpus_dir, store=store, provider=provider, settings=settings)
    assert result["failed"] == 1
    assert provider.calls == []
    assert store.counts() == (0, 0)


async def test_empty_directory_retires_deleted_source_but_preserves_evidence(settings, store, provider, source, ingested):
    record = active_records(store)[0]
    source.unlink()
    result = await ingest_directory(settings.corpus_dir, store=store, provider=provider, settings=settings)
    assert result["retired"] == 1
    assert result["items"][0]["document_id"] == ingested["document_id"]
    assert store.counts() == (0, 0)
    assert Path(record["original_path"]).is_file()
    with store.connect() as conn:
        assert conn.execute("SELECT count(*) FROM chunks WHERE document_id=?", (ingested["document_id"],)).fetchone()[0] > 0


async def test_rename_activates_new_path_and_retires_old_path(settings, store, provider, source, ingested):
    renamed = source.with_name("renamed.md")
    source.rename(renamed)
    result = await ingest_directory(settings.corpus_dir, store=store, provider=provider, settings=settings)
    assert result["ingested"] == result["retired"] == 1
    assert [row["source_path"] for row in active_records(store)] == [str(renamed.resolve())]
    with store.connect() as conn:
        assert conn.execute("SELECT count(*) FROM documents").fetchone()[0] == 2


async def test_directory_sync_does_not_retire_other_corpus_or_failed_existing_source(settings, store, provider, source, ingested, tmp_path):
    other = tmp_path / "separate" / "other.md"
    other.parent.mkdir()
    other.write_text(SOURCE_TEXT, encoding="utf-8")
    other_result = await ingest_file(other, store=store, provider=provider, settings=settings)
    other.unlink()
    source.write_bytes(b"\xff invalid utf-8 source")
    result = await ingest_directory(settings.corpus_dir, store=store, provider=provider, settings=settings)
    assert result["failed"] == 1
    assert result["retired"] == 0
    assert {row["id"] for row in active_records(store)} == {ingested["document_id"], other_result["document_id"]}


async def test_discovery_failure_does_not_retire_missing_documents(settings, store, provider, source, ingested, monkeypatch):
    import app.ingestion as module
    source.unlink()

    def failing_walk(directory, *, followlinks, onerror):
        onerror(PermissionError("Synthetic inaccessible corpus child"))
        yield  # This function remains a walk-like iterator.

    monkeypatch.setattr(module.os, "walk", failing_walk)
    with pytest.raises(PermissionError, match="inaccessible corpus child"):
        await ingest_directory(settings.corpus_dir, store=store, provider=provider, settings=settings)
    assert [row["id"] for row in active_records(store)] == [ingested["document_id"]]


async def test_ingest_one_file_does_not_retire_missing_siblings(settings, store, provider, source, ingested):
    second = source.with_name("second.md")
    second.write_text(SOURCE_TEXT, encoding="utf-8")
    source.unlink()
    await ingest_file(second, store=store, provider=provider, settings=settings)
    assert len(active_records(store)) == 2


async def test_external_source_link_is_skipped_visibly(settings, store, provider, tmp_path):
    settings.corpus_dir.mkdir()
    outside = tmp_path / "outside.md"
    outside.write_text(SOURCE_TEXT, encoding="utf-8")
    link = settings.corpus_dir / "linked.md"
    try:
        link.symlink_to(outside)
    except OSError as error:
        pytest.skip(f"This Windows account cannot create symlinks: {error.winerror}")
    result = await ingest_directory(settings.corpus_dir, store=store, provider=provider, settings=settings)
    assert result["skipped"] == 1
    assert result["items"][0]["source_path"] == "linked.md"
    assert provider.calls == []


@pytest.mark.parametrize("corrupted", ["not json", "null", "{}", '"vector"', "[]", "[1, 0]", "[0, 0, 0]", "[true, 0, 0]", '["1", 0, 0]', "[NaN, 0, 1]", "[Infinity, 0, 1]"])
async def test_corrupted_embedding_cache_is_evicted_and_refetched(settings, store, provider, corrupted):
    text = "A cached query"
    key = hashlib.sha256((settings.openrouter_embedding_model + "\0" + str(settings.openrouter_embedding_dimensions) + "\0" + text).encode()).hexdigest()
    with store.connect() as conn:
        conn.execute("INSERT INTO embedding_cache VALUES(?,?)", (key, corrupted))
    result = await cached_embeddings(store, provider, settings, [text])
    assert result == [[1.0, 0.0, 0.0]]
    assert provider.calls == [("embeddings", [text])]
    with store.connect() as conn:
        assert json.loads(conn.execute("SELECT vector FROM embedding_cache WHERE cache_key=?", (key,)).fetchone()[0]) == result[0]
        assert json.loads(conn.execute("SELECT payload FROM events WHERE kind='embedding_cache_repair'").fetchone()[0]) == {"evicted_entries": 1}


@pytest.mark.parametrize("vectors", [None, {}, [[0, 0, 0]], [[1, 0]], [[float("nan"), 0, 1]], [[True, 0, 1]]])
async def test_invalid_provider_vectors_never_enter_embedding_cache(settings, store, vectors):
    class InvalidEmbeddingProvider(FakeProvider):
        async def embeddings(self, texts):
            return vectors

    with pytest.raises(ProviderError) as error:
        await cached_embeddings(store, InvalidEmbeddingProvider(), settings, ["query"])
    assert error.value.category == "invalid_embedding"
    with store.connect() as conn:
        assert conn.execute("SELECT count(*) FROM embedding_cache").fetchone()[0] == 0
